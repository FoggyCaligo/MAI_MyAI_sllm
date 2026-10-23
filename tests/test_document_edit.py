"""Structured document edits and owner-only exposure."""
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter

from mai.tools.document_edit import (
    CsvSetCell, DocxSetParagraph, DocxSetTableCell, DocumentEditInput,
    PdfDeletePage, PdfRotatePage, PptxSetShapeText, XlsxSetCell, document_edit,
)
from mai.tools.local import register_local_pc_tools, register_readonly_local_tools
from mai.tools.registry import ToolRegistry


def test_docx_paragraph_edit_and_copy(tmp_path: Path) -> None:
    source = tmp_path / "input.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("Before")
    doc.save(source)
    result = document_edit(path=str(source), output_path=str(output),
                           action=DocxSetParagraph(operation="docx_set_paragraph", paragraph=1, text="After"))
    assert result["operation"] == "docx_set_paragraph"
    assert Document(source).paragraphs[0].text == "Before"
    assert Document(output).paragraphs[0].text == "After"


def test_docx_table_cell_edit(tmp_path: Path) -> None:
    source = tmp_path / "table.docx"
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "old"
    doc.save(source)
    document_edit(path=str(source), action=DocxSetTableCell(
        operation="docx_set_table_cell", table=1, row=1, column=1, text="new"))
    assert Document(source).tables[0].cell(0, 0).text == "new"


def test_xlsx_cell_edit_and_nonexistent_sheet_preserves_source(tmp_path: Path) -> None:
    source = tmp_path / "cells.xlsx"
    book = Workbook()
    book.active.title = "Data"
    book.active["A1"] = "old"
    book.save(source)
    document_edit(path=str(source), action=XlsxSetCell(
        operation="xlsx_set_cell", sheet="Data", cell="A1", value="new"))
    with load_workbook(source) as result:
        assert result["Data"]["A1"].value == "new"
    before = source.read_bytes()
    with pytest.raises(KeyError):
        document_edit(path=str(source), action=XlsxSetCell(
            operation="xlsx_set_cell", sheet="Missing", cell="A1", value="x"))
    assert source.read_bytes() == before


def test_pptx_shape_edit(tmp_path: Path) -> None:
    source = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = "Before"
    presentation.save(source)
    document_edit(path=str(source), action=PptxSetShapeText(
        operation="pptx_set_shape_text", slide=1, shape=1, text="After"))
    assert Presentation(source).slides[0].shapes[0].text == "After"


def test_pdf_rotate_and_delete_pages(tmp_path: Path) -> None:
    source = tmp_path / "pages.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=200)
    writer.add_blank_page(width=100, height=200)
    with source.open("wb") as handle:
        writer.write(handle)
    document_edit(path=str(source), action=PdfRotatePage(
        operation="pdf_rotate_page", page=1, degrees=90))
    assert PdfReader(source).pages[0].get("/Rotate") == 90
    document_edit(path=str(source), action=PdfDeletePage(
        operation="pdf_delete_page", page=2))
    assert len(PdfReader(source).pages) == 1
    before = source.read_bytes()
    with pytest.raises(ValueError, match="last PDF page"):
        document_edit(path=str(source), action=PdfDeletePage(
            operation="pdf_delete_page", page=1))
    assert source.read_bytes() == before


def test_csv_cell_edit(tmp_path: Path) -> None:
    source = tmp_path / "data.csv"
    source.write_text("a,b\nc,d\n", encoding="utf-8")
    document_edit(path=str(source), encoding="utf-8", action=CsvSetCell(
        operation="csv_set_cell", row=2, column=2, value="replaced"))
    assert source.read_text(encoding="utf-8").splitlines()[1] == "c,replaced"


def test_mismatched_file_extension_rejected_without_writing(tmp_path: Path) -> None:
    source = tmp_path / "content.csv"
    source.write_text("a,b\n")
    with pytest.raises(ValueError, match="requires input"):
        document_edit(path=str(source), action=DocxSetParagraph(
            operation="docx_set_paragraph", paragraph=1, text="x"))
    assert source.read_text() == "a,b\n"


def test_structural_action_schema_rejects_extra_fields() -> None:
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        DocumentEditInput.model_validate({
            "path": "test.xlsx",
            "action": {"operation": "xlsx_set_cell", "sheet": "Sheet", "cell": "A1",
                       "value": "hi", "unexpected": True},
        })


def test_document_edit_is_owner_only(tmp_path: Path) -> None:
    owner = ToolRegistry()
    register_local_pc_tools(owner, cwd=tmp_path)
    assert "document_edit" in owner.names()
    trial = ToolRegistry()
    register_readonly_local_tools(trial, cwd=tmp_path, upload_root=tmp_path / "uploads")
    assert "document_edit" not in trial.names()
