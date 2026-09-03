from __future__ import annotations

from docx import Document
from openpyxl import Workbook

from mai.tools.documents import document_read, register_document_tools
from mai.tools.filesystem import file_read, register_filesystem_read_tools
from mai.tools.images import register_image_tools
from mai.tools.registry import ToolRegistry


def test_file_read_handles_docx_xlsx_and_csv(tmp_path):
    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("hello document")
    doc.save(docx_path)
    docx_result = file_read(path=str(docx_path))
    assert docx_result["extension"] == ".docx"
    assert "hello document" in docx_result["content"]

    xlsx_path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "hello sheet"
    workbook.save(xlsx_path)
    xlsx_result = file_read(path=str(xlsx_path))
    assert xlsx_result["extension"] == ".xlsx"
    assert "hello sheet" in xlsx_result["content"]

    csv_path = tmp_path / "sample.csv"
    csv_path.write_text("name,value\nalpha,10\nbeta,20\n", encoding="utf-8-sig")
    csv_result = file_read(path=str(csv_path))
    assert csv_result["extension"] == ".csv"
    assert csv_result["details"]["row_count"] == 3
    assert "alpha\t10" in csv_result["content"]


def test_file_read_csv_accepts_explicit_encoding(tmp_path):
    csv_path = tmp_path / "korean.csv"
    csv_path.write_text("이름,값\n테스트,10\n", encoding="cp949")
    result = file_read(path=str(csv_path), encoding="cp949")
    assert result["details"]["encoding"] == "cp949"
    assert "테스트\t10" in result["content"]


def test_document_read_remains_available_for_internal_callers(tmp_path):
    csv_path = tmp_path / "sample.csv"
    csv_path.write_text("name,value\nalpha,10\n", encoding="utf-8-sig")
    result = document_read(path=str(csv_path))
    assert result["extension"] == ".csv"
    assert "alpha\t10" in result["text"]


def test_registry_exposes_one_file_read_interface_for_documents(tmp_path):
    registry = ToolRegistry()
    register_document_tools(registry, cwd=tmp_path)
    register_filesystem_read_tools(registry, cwd=tmp_path)
    register_image_tools(registry, model="vision-test", host="http://127.0.0.1:11434", cwd=tmp_path)
    names = set(registry.names())
    assert "file_read" in names
    assert "document_read" not in names
    assert "image_analyze" in names
