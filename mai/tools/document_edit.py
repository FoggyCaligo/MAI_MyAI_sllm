"""Structured, transactional document edits using existing MAI dependencies.

An edit is one explicit typed operation. No interpretation of free-text
instructions or file-extension heuristics for selecting an operation.
"""
from __future__ import annotations

import csv
import os
import tempfile
from pathlib import Path
from typing import Annotated, Any, Literal, Union

from pydantic import BaseModel, ConfigDict, Field

from .registry import ToolRegistry


class _Strict(BaseModel):
    model_config = ConfigDict(extra="forbid")


class DocxSetParagraph(_Strict):
    operation: Literal["docx_set_paragraph"]
    paragraph: int = Field(ge=1, description="1-based index among top-level document paragraphs")
    text: str


class DocxSetTableCell(_Strict):
    operation: Literal["docx_set_table_cell"]
    table: int = Field(ge=1)
    row: int = Field(ge=1)
    column: int = Field(ge=1)
    text: str


class XlsxSetCell(_Strict):
    operation: Literal["xlsx_set_cell"]
    sheet: str = Field(min_length=1)
    cell: str = Field(min_length=1, description="A1-style cell reference")
    value: str | int | float | bool | None


class PptxSetShapeText(_Strict):
    operation: Literal["pptx_set_shape_text"]
    slide: int = Field(ge=1)
    shape: int = Field(ge=1, description="1-based index in slide.shapes")
    text: str


class PdfRotatePage(_Strict):
    operation: Literal["pdf_rotate_page"]
    page: int = Field(ge=1)
    degrees: Literal[90, 180, 270]


class PdfDeletePage(_Strict):
    operation: Literal["pdf_delete_page"]
    page: int = Field(ge=1)


class CsvSetCell(_Strict):
    operation: Literal["csv_set_cell"]
    row: int = Field(ge=1)
    column: int = Field(ge=1)
    value: str


DocumentAction = Annotated[
    Union[DocxSetParagraph, DocxSetTableCell, XlsxSetCell, PptxSetShapeText,
          PdfRotatePage, PdfDeletePage, CsvSetCell],
    Field(discriminator="operation"),
]


class DocumentEditInput(_Strict):
    path: str = Field(min_length=1)
    action: DocumentAction
    output_path: str | None = Field(default=None, description="Optional different path; omitted means modify the original file")
    encoding: str = "utf-8-sig"


_FORMAT_BY_OPERATION: dict[str, str] = {
    "docx_set_paragraph": ".docx",
    "docx_set_table_cell": ".docx",
    "xlsx_set_cell": ".xlsx",
    "pptx_set_shape_text": ".pptx",
    "pdf_rotate_page": ".pdf",
    "pdf_delete_page": ".pdf",
    "csv_set_cell": ".csv",
}


def _resolve(path: str, cwd: str | Path | None) -> Path:
    candidate = Path(path).expanduser()
    if not candidate.is_absolute():
        candidate = Path(cwd or os.getcwd()) / candidate
    return candidate.resolve(strict=False)


def _at(items: Any, index: int, label: str) -> Any:
    if index > len(items):
        raise IndexError(f"{label} index {index} exceeds available count {len(items)}")
    return items[index - 1]


def document_edit(
    *, path: str, action: DocumentAction, output_path: str | None = None,
    encoding: str = "utf-8-sig", cwd: str | Path | None = None,
) -> dict[str, Any]:
    source = _resolve(path, cwd)
    if not source.exists():
        raise FileNotFoundError(str(source))
    if not source.is_file():
        raise IsADirectoryError(str(source))
    destination = _resolve(output_path, cwd) if output_path is not None else source
    expected_extension = _FORMAT_BY_OPERATION[action.operation]
    if source.suffix.lower() != expected_extension or destination.suffix.lower() != expected_extension:
        raise ValueError(f"{action.operation} requires input and output {expected_extension} files")
    if not destination.parent.is_dir():
        raise FileNotFoundError(str(destination.parent))

    if expected_extension == ".docx":
        from docx import Document
        doc = Document(source)
        if isinstance(action, DocxSetParagraph):
            # Setting paragraph.text intentionally replaces its inline runs and
            # may remove their character-level formatting; the paragraph remains.
            _at(doc.paragraphs, action.paragraph, "paragraph").text = action.text
        elif isinstance(action, DocxSetTableCell):
            table = _at(doc.tables, action.table, "table")
            row = _at(table.rows, action.row, "row")
            _at(row.cells, action.column, "column").text = action.text
        else:
            raise TypeError("invalid DOCX action")
        save = doc.save
    elif expected_extension == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(source)
        try:
            if not isinstance(action, XlsxSetCell):
                raise TypeError("invalid XLSX action")
            if action.sheet not in workbook.sheetnames:
                raise KeyError(f"worksheet not found: {action.sheet}")
            worksheet = workbook[action.sheet]
            worksheet[action.cell].value = action.value
        except BaseException:
            workbook.close()
            raise
        save = workbook.save
    elif expected_extension == ".pptx":
        from pptx import Presentation
        presentation = Presentation(source)
        if not isinstance(action, PptxSetShapeText):
            raise TypeError("invalid PPTX action")
        slide = _at(presentation.slides, action.slide, "slide")
        shape = _at(slide.shapes, action.shape, "shape")
        if not shape.has_text_frame:
            raise TypeError("selected shape has no text frame")
        # shape.text replaces text runs and their formatting.
        shape.text = action.text
        save = presentation.save
    elif expected_extension == ".pdf":
        from pypdf import PdfReader, PdfWriter
        reader = PdfReader(source)
        if len(reader.pages) == 0:
            raise ValueError("PDF has no pages")
        writer = PdfWriter()
        if isinstance(action, PdfRotatePage):
            _at(reader.pages, action.page, "page").rotate(action.degrees)
            for page in reader.pages:
                writer.add_page(page)
        elif isinstance(action, PdfDeletePage):
            _at(reader.pages, action.page, "page")
            if len(reader.pages) == 1:
                raise ValueError("cannot remove the last PDF page")
            for idx, page in enumerate(reader.pages, start=1):
                if idx != action.page:
                    writer.add_page(page)
        else:
            raise TypeError("invalid PDF action")
        save = writer.write
    elif expected_extension == ".csv":
        if not isinstance(action, CsvSetCell):
            raise TypeError("invalid CSV action")
        with source.open("r", encoding=encoding, newline="") as handle:
            rows = list(csv.reader(handle))
        row = _at(rows, action.row, "row")
        _at(row, action.column, "column")
        row[action.column - 1] = action.value
        def save(destination_file: str | Path) -> None:
            with Path(destination_file).open("w", encoding=encoding, newline="") as handle:
                csv.writer(handle).writerows(rows)
    else:
        raise ValueError(f"unsupported document extension: {expected_extension}")

    # Save to a sibling temporary file; a failed edit never partially overwrites
    # the original. Replacing the destination happens only after saving succeeds.
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(dir=destination.parent, suffix=expected_extension,
                                         prefix=".mai-document-", delete=False) as tmp:
            temporary_path = Path(tmp.name)
        save(temporary_path)
        os.replace(temporary_path, destination)
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        if expected_extension == ".xlsx":
            workbook.close()

    return {
        "path": str(destination),
        "source_path": str(source),
        "operation": action.operation,
        "bytes": destination.stat().st_size,
    }


def register_document_edit_tool(
    registry: ToolRegistry, *, cwd: str | Path | None = None,
    timeout_seconds: float | None = 60,
) -> None:
    def handler(**kwargs: Any) -> dict[str, Any]:
        # ToolRegistry passes Pydantic's model_dump(), so restore the
        # discriminated action model before dispatching to typed handlers.
        validated = DocumentEditInput.model_validate(kwargs, strict=True)
        return document_edit(
            path=validated.path,
            action=validated.action,
            output_path=validated.output_path,
            encoding=validated.encoding,
            cwd=cwd,
        )

    registry.add(
        name="document_edit",
        description=(
            "Edit an existing DOCX, XLSX, PPTX, PDF, or CSV file using a single explicit typed action. "
            "DOCX: set a top-level paragraph or table cell; XLSX: set an A1-style cell value; "
            "PPTX: set shape text; PDF: rotate or delete a page (not arbitrary PDF text editing); "
            "CSV: set an existing cell. All indices are 1-based; missing targets fail. "
            "Paragraph and shape text replacement can reset inline formatting. "
            "Save atomically; output_path may preserve the source unchanged. "
            "No conversion or file creation; select a documented operation exactly."
        ),
        input_model=DocumentEditInput,
        handler=handler,
        timeout_seconds=timeout_seconds,
        category="document",
    )
